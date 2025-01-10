import os
import requests
from bs4 import BeautifulSoup
from docx import Document


def Get_All_Href(url):
    """
    如果需要从指定 URL 中抓取 <nav> 下所有链接，可使用此函数。
    返回完整链接列表。
    """
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/114.0.0.0 Safari/537.36"
        )
    }
    response = requests.get(url, headers=headers)
    response.raise_for_status()
    html_content = response.text

    soup = BeautifulSoup(html_content, 'html.parser')
    nav = soup.find('nav', class_='sidenav sidenav-large custom-text')
    base_url = 'https://developer.apple.com'

    if not nav:
        print("警告：未找到 <nav>，可能结构已变动。")
        return []

    hrefs = []
    for a in nav.find_all('a', href=True):
        # 拼接成完整链接
        full_link = a['href']
        if not full_link.startswith('http'):
            full_link = base_url + full_link
        hrefs.append(full_link)

    return hrefs


def parse_list(list_tag, doc, ordered=False):
    """
    解析 <ol> 或 <ul> 下的 <li> 列表项，并写入到 Word 文档中。
    ordered=True 表示有序列表（自动加数字序号），
    ordered=False 表示无序列表（使用 Bullet 样式）。
    """
    list_counter = 1
    # 如果有嵌套列表，需要再加递归处理。这里只演示一层。
    for li in list_tag.find_all('li', recursive=False):
        text = li.get_text(strip=True)
        if not text:
            continue
        if ordered:
            # 有序列表：1.、2.、3. ...
            doc.add_paragraph(f"{list_counter}. {text}")
            list_counter += 1
        else:
            # 无序列表：用 Word 的 Bullet 样式
            doc.add_paragraph(text, style='List Bullet')


def parse_table(table, doc, table_counter):
    """
    解析表格，插入到 Word 文档中。
    1) 尝试自动获取表格标题（从父级 Subhead 或 aria-label 或者默认“表格 x”）。
    2) 将表格每一行写入 Word 表格。
    返回更新后的 table_counter。
    """
    table_title = None

    # 先找父级 <div class="Subhead"> 中的 <h2>
    subhead_div = table.find_parent('div', class_='Subhead')
    if subhead_div:
        h2 = subhead_div.find('h2')
        if h2:
            table_title = h2.get_text(strip=True)

    # 如果还没找到标题，就看 <table> 是否有 aria-label
    if not table_title:
        aria_label = table.get('aria-label')
        if aria_label:
            table_title = aria_label.strip()

    # 如果还没有，就用“表格 x”
    if not table_title:
        table_title = f"表格 {table_counter}"

    # 在 Word 中插入表格标题
    doc.add_heading(table_title, level=2)
    table_counter += 1

    # 开始解析表格数据
    rows = table.find_all('tr')
    if not rows:
        return table_counter

    # 取第一行，看看有多少列
    first_row_cols = rows[0].find_all(['th', 'td'])
    if not first_row_cols:
        return table_counter
    num_cols = len(first_row_cols)

    # 在 Word 中创建对应列数的表格
    word_table = doc.add_table(rows=1, cols=num_cols)
    word_table.style = 'Table Grid'

    # 处理表头
    hdr_cells = word_table.rows[0].cells
    for i, cell in enumerate(first_row_cols):
        hdr_cells[i].text = cell.get_text(strip=True)

    # 处理后面几行
    for row in rows[1:]:
        cells = row.find_all(['td', 'th'])
        if len(cells) != num_cols:
            # 如果某行列数不一致，可以选择跳过或自行处理
            continue
        row_cells = word_table.add_row().cells
        for j, cell in enumerate(cells):
            row_cells[j].text = cell.get_text(strip=True)

    return table_counter


def parse_element(element, doc, table_counter):
    """
    递归解析单个 DOM element，根据标签类型写入文档。
    如果 element 下还有子节点，需要继续深入（如 <div>、<section> 等）。
    返回更新后的 table_counter。
    """
    if element.name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
        # 标题
        level = int(element.name[-1])  # h1->1, h2->2 ...
        doc.add_heading(element.get_text(strip=True), level=level)

    elif element.name == 'p':
        # 段落
        p_text = element.get_text(strip=True)
        if p_text:
            doc.add_paragraph(p_text)

    elif element.name == 'ol':
        # 有序列表
        parse_list(element, doc, ordered=True)

    elif element.name == 'ul':
        # 无序列表
        parse_list(element, doc, ordered=False)

    elif element.name == 'table':
        # 直接是表格
        table_counter = parse_table(element, doc, table_counter)

    elif element.name == 'div' and 'table-wrapper' in (element.get('class') or []):
        # 有些页面会用 <div class="table-wrapper"> 包裹 <table>
        maybe_table = element.find('table')
        if maybe_table:
            table_counter = parse_table(maybe_table, doc, table_counter)
        else:
            # 如果没有 table，就把里面的文本当成普通段落
            text_in_div = element.get_text(strip=True)
            if text_in_div:
                doc.add_paragraph(text_in_div)

    else:
        # 对于其他标签（如 article, section, div 等），
        # 可以递归往下找它的子节点
        for child in element.children:
            if child.name:  # 排除纯文本等
                table_counter = parse_element(child, doc, table_counter)

    return table_counter


def scrape_and_save_to_word(url, output_directory):
    """
    针对单个 URL:
      1) 请求并解析 <div class="main-content"> 内容
      2) 按顺序递归处理所有子节点
      3) 将最终结果保存到 Word 文档
    文件名使用最后一个路径片段 + .docx；若为空则用 index.docx
    """
    headers = {"User-Agent": "Mozilla/5.0"}
    try:
        response = requests.get(url, headers=headers)
        response.raise_for_status()
    except requests.RequestException as e:
        print(f"请求失败: {url} -> {e}")
        return

    soup = BeautifulSoup(response.text, 'html.parser')
    main_content = soup.find('div', class_='main-content')
    if not main_content:
        print(f"未找到 main-content 区域: {url}")
        return

    doc = Document()
    table_counter = 1

    # 依次解析 main_content 的子节点
    for child in main_content.children:
        if child.name:
            table_counter = parse_element(child, doc, table_counter)

    # 根据 url 的最后一个路径段作为文件名
    filename = url.rstrip('/').split('/')[-1]
    if not filename:
        filename = "index"
    filename += ".docx"

    # 最终保存路径
    output_path = os.path.join(output_directory, filename)
    doc.save(output_path)
    print(f"已保存: {output_path}")


def main():
    # 1) 先从指定页面抓取所有导航链接
    base_url = "https://developer.apple.com/cn/help/app-store-connect/get-started/app-store-connect-homepage"
    all_hrefs = Get_All_Href(base_url)

    # 2) 准备输出目录
    output_directory = r"D:\Nick\work\AI_对话\pacongwenjian\test"
    os.makedirs(output_directory, exist_ok=True)

    # 3) 逐个链接爬取并保存为 Word
    for href in all_hrefs:
        scrape_and_save_to_word(href, output_directory)


if __name__ == "__main__":
    main()
